import sys
import openpyxl
import os

# Set output encoding to UTF-8
sys.stdout.reconfigure(encoding='utf-8')

def inspect_xlsx():
    cwd = r"d:\Quản lý vận hành DYC"
    files = [f for f in os.listdir(cwd) if f.endswith(".xlsx")]
    for file in files:
        file_path = os.path.join(cwd, file)
        wb = openpyxl.load_workbook(file_path, read_only=True)
        print(f"\nFile: {file}")
        print(f"Sheets: {wb.sheetnames}")
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            # Convert to list to avoid sheet closed errors
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(row)
            print(f"Sheet: {sheet_name}, Max Rows: {len(rows)}")
            for i in range(min(12, len(rows))):
                print(f"  Row {i+1}: {rows[i][:8]}")

if __name__ == "__main__":
    inspect_xlsx()
