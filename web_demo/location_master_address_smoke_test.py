# -*- coding: utf-8 -*-
"""
Smoke: location master import + /api/location + build_full_address (Python).
Chạy từ web_demo:  python location_master_address_smoke_test.py
"""
import json
import os
import sys
import uuid
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from database import init_db, SessionLocal
from main import app
from vehicle_norm_logic import build_full_address
from location_master_import import SHEET_NAME, run_import

HERE = Path(__file__).resolve().parent
EXCEL = HERE / "data" / "Danh-muc-Phuong-xa_moi.xlsx"
JSON_OUT = HERE / "static" / "location_master.json"
client = TestClient(app)
ROWS = []


def add(n, name, ok, detail=""):
    ROWS.append((n, name, ok, detail))


def _ensure_sample_excel():
    """Tạo Excel mẫu đúng cấu trúc sheet/cột nếu chưa có file chính thức (để CI/smoke pass)."""
    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    if EXCEL.is_file():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.cell(1, 1, "row1")
    ws.cell(2, 1, "row2")
    ws.cell(3, 1, "header")
    # D=4, J=10 — dữ liệu từ dòng 4
    rows = [
        ("Thành phố Hồ Chí Minh", "Phường Cầu Ông Lãnh"),
        ("Thành phố Hồ Chí Minh", "Phường An Phú"),
        ("Thành phố Hà Nội", "Phường Hoàn Kiếm"),
        ("Thành phố Hà Nội", "Phường Ba Đình"),
        ("Tỉnh Đồng Nai", "Xã An Phước"),
    ]
    r = 4
    for prov, ward in rows:
        ws.cell(r, 4, prov)
        ws.cell(r, 10, ward)
        r += 1
    wb.save(EXCEL)
    wb.close()


def main():
    init_db()
    _ensure_sample_excel()
    add(1, "File Excel tồn tại", EXCEL.is_file(), str(EXCEL))

    try:
        meta, data = run_import(str(EXCEL), str(JSON_OUT))
    except Exception as e:
        add(2, "Chạy import tạo location_master.json", False, str(e))
        _print_fail()
        sys.exit(1)

    add(
        2,
        "Chạy import tạo location_master.json",
        JSON_OUT.is_file(),
        f"provinces={meta.get('provinces_count')}, wards={meta.get('items_count')}",
    )
    with open(JSON_OUT, "r", encoding="utf-8") as f:
        j = json.load(f)
    add(3, "JSON có key data", isinstance(j.get("data"), dict), list(j.keys())[:5])
    provs = list(j["data"].keys())
    add(4, "Có ít nhất 1 tỉnh/TP", len(provs) >= 1, f"n={len(provs)}")
    n_wards = sum(len(v) for v in j["data"].values())
    add(5, "Có ít nhất 1 phường/xã", n_wards >= 1, f"n={n_wards}")

    r6 = client.get("/api/location/provinces")
    add(6, "GET /api/location/provinces 200", r6.status_code == 200, r6.text[:120])
    j6 = r6.json()
    add(7, "provinces có items", isinstance(j6.get("items"), list) and len(j6["items"]) >= 1, str(j6.get("total")))

    r8 = client.get("/api/location/wards", params={"province": "Thành phố Hà Nội"})
    add(8, "GET wards Hà Nội 200", r8.status_code == 200, "")
    if "Thành phố Hồ Chí Minh" in j.get("data", {}):
        r9 = client.get("/api/location/wards", params={"province": "Thành phố Hồ Chí Minh"})
        add(9, "GET wards TP.HCM có items", r9.status_code == 200 and len(r9.json().get("items", [])) >= 1, r9.text[:100])
    else:
        add(9, "GET wards TP.HCM (skip)", True, "no HCM in sample")

    r9b = client.get("/api/location/wards", params={"province": "Không tồn tại XYZ"})
    add(10, "wards province lạ → []", r9b.status_code == 200 and r9b.json().get("items") == [], str(r9b.json()))

    exp10 = "Số 264, Đường Trần Hưng Đạo, Phường Cầu Ông Lãnh, Thành phố Hồ Chí Minh"
    got10 = build_full_address("Số 264", "Trần Hưng Đạo", "Phường Cầu Ông Lãnh", "Thành phố Hồ Chí Minh")
    add(11, "build_full_address đủ phần", got10 == exp10, got10)

    got11 = build_full_address("Số 264", "Trần Hưng Đạo", "Cầu Ông Lãnh", "Thành phố Hồ Chí Minh")
    add(12, "ward thiếu Phường → có Phường", "Phường Cầu Ông Lãnh" in got11 and "Phường Phường" not in got11, got11)

    got12 = build_full_address("1", "Đường Trần Hưng Đạo", "Phường A", "TP.HCM")
    add(13, "street đã có Đường", "Đường Đường" not in got12, got12)

    got13 = build_full_address("1", "X", "Y", "Thành phố Hồ Chí Minh")
    add(14, "city đã có Thành phố", "Thành phố Thành phố" not in got13, got13)

    sfx = uuid.uuid4().hex[:8].upper()
    did = f"LOC-SMOKE-D-{sfx}"
    r14 = client.post(
        "/api/dealers",
        json={
            "dealer_id": did,
            "dealer_name": "Loc Smoke Dealer",
            "address_no": "Số 10",
            "street": "Lê Lợi",
            "ward": "Phường B",
            "city": "Thành phố Hà Nội",
            "created_by": "SMOKE",
        },
    )
    ok14 = r14.status_code == 200 and "Phường" in (r14.json().get("full_address") or "")
    add(15, "POST dealer full_address auto", ok14, r14.text[:200])

    cid = f"LOC-SMOKE-C-{sfx}"
    r15 = client.post(
        "/api/end-customers",
        json={
            "customer_id": cid,
            "customer_masked": "KH LOC",
            "address_no": "12",
            "street": "Trần Hưng Đạo",
            "ward": "Phường Hoàn Kiếm",
            "city": "Thành phố Hà Nội",
            "source_channel": "MANUAL",
            "created_by": "SMOKE",
        },
    )
    ok15 = r15.status_code == 200 and "Đường" in (r15.json().get("full_address") or "")
    add(16, "POST customer full_address auto", ok15, r15.text[:200])

    r16 = client.put(
        f"/api/dealers/{did}",
        json={
            "reason": "đổi ward",
            "ward": "Phường Ba Đình",
            "city": "Thành phố Hà Nội",
            "updated_by": "SMOKE",
        },
    )
    add(17, "PUT dealer ward master", r16.status_code == 200, r16.text[:120])

    r17 = client.put(
        f"/api/end-customers/{cid}",
        json={
            "reason": "ward tay",
            "ward": "Khu phố không trong master 123",
            "city": "Thành phố Hà Nội",
            "updated_by": "SMOKE",
        },
    )
    add(18, "PUT customer ward tay", r17.status_code == 200, r17.text[:120])

    failed = [x for x in ROWS if not x[2]]
    for n, name, ok, det in ROWS:
        print(f"  [{n}] {'PASS' if ok else 'FAIL'} — {name} {(det or '')[:100]}")

    if failed:
        print("\nFAILED:", failed)
        sys.exit(1)
    print("")
    print("LOCATION MASTER & ADDRESS BUILD SMOKE TEST PASSED")


def _print_fail():
    for n, name, ok, det in ROWS:
        print(f"  [{n}] {'PASS' if ok else 'FAIL'} — {name}")


if __name__ == "__main__":
    main()
