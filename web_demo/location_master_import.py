# -*- coding: utf-8 -*-
"""
Import danh mục Tỉnh/TP (cột C) và Phường/Xã (cột I) từ CSV → static/location_master.json.

Chạy từ thư mục web_demo:
  python location_master_import.py

File nguồn mặc định: data/Danh-muc-Phuong-xa_moi.csv
  (hoặc thư mục cha dự án: ../Danh-muc-Phuong-xa_moi.csv)

CSV: dòng 1 là header, dữ liệu từ dòng 2. Encoding UTF-8 (có BOM cũng được).

Tùy chọn Excel cũ: truyền --input path/to/file.xlsx (sheet `1.DM Phường xã mới`, cột D/J, dữ liệu từ dòng 4).
"""
from __future__ import annotations

import argparse
import csv
import json
import locale
import os
import re
import sys
from collections import defaultdict
from typing import Dict, List, Optional, Tuple

# --- CSV (mặc định): cột C = index 2, cột I = index 8 (0-based) ---
CSV_COL_PROVINCE = 2  # C — "Tên tỉnh/TP mới"
CSV_COL_WARD = 8  # I — "Tên Phường/Xã mới"
CSV_DATA_START_ROW = 2  # dòng 1 header, dòng 2+ dữ liệu

# --- Excel (legacy) ---
SHEET_NAME = "1.DM Phường xã mới"
XLS_COL_PROVINCE = 4  # D
XLS_COL_WARD = 10  # J
XLS_HEADER_ROW = 3
XLS_DATA_START_ROW = 4


def _base_dir() -> str:
    return os.path.dirname(os.path.abspath(__file__))


def default_csv_path() -> str:
    return os.path.join(_base_dir(), "data", "Danh-muc-Phuong-xa_moi.csv")


def default_output_path() -> str:
    return os.path.join(_base_dir(), "static", "location_master.json")


def _normalize_cell(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _sort_vi(strings: List[str]) -> List[str]:
    """Sắp xếp ổn định; thử locale vi nếu có rồi khôi phục LC_COLLATE."""
    lst = list(strings)
    old_loc = None
    try:
        old_loc = locale.setlocale(locale.LC_COLLATE)
    except Exception:
        pass
    try:
        for loc in ("vi_VN.UTF-8", "Vietnamese_Vietnam.utf8", "en_US.UTF-8"):
            try:
                locale.setlocale(locale.LC_COLLATE, loc)
                return sorted(lst, key=locale.strxfrm)
            except Exception:
                continue
        return sorted(lst, key=lambda x: (x.lower(), x))
    finally:
        if old_loc:
            try:
                locale.setlocale(locale.LC_COLLATE, old_loc)
            except Exception:
                pass


def _resolve_input_path(input_path: Optional[str]) -> str:
    if input_path and os.path.isfile(input_path):
        return input_path
    p = default_csv_path()
    if os.path.isfile(p):
        return p
    alt = os.path.join(os.path.dirname(_base_dir()), "Danh-muc-Phuong-xa_moi.csv")
    if os.path.isfile(alt):
        print(f"INFO: dùng file CSV thay thế: {alt}")
        return alt
    alt_xlsx = os.path.join(_base_dir(), "data", "Danh-muc-Phuong-xa_moi.xlsx")
    if os.path.isfile(alt_xlsx):
        print(f"INFO: không thấy CSV, dùng Excel: {alt_xlsx}")
        return alt_xlsx
    alt_xlsx2 = os.path.join(os.path.dirname(_base_dir()), "Danh-muc-Phuong-xa_moi.xlsx")
    if os.path.isfile(alt_xlsx2):
        print(f"INFO: không thấy CSV, dùng Excel: {alt_xlsx2}")
        return alt_xlsx2
    raise FileNotFoundError(
        "Không tìm thấy danh mục. Đặt một trong các file:\n"
        f"  {default_csv_path()}\n"
        f"  {alt}\n"
        f"  (hoặc .xlsx legacy trong data/ hoặc thư mục cha dự án.)"
    )


def _import_from_csv(csv_path: str) -> Tuple[dict, Dict[str, List[str]]]:
    by_province: Dict[str, set] = defaultdict(set)
    total_rows = 0
    base = os.path.basename(csv_path)
    try:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                raise ValueError("CSV rỗng hoặc thiếu header")
            max_idx = max(CSV_COL_PROVINCE, CSV_COL_WARD)
            for row in reader:
                if not row or all(not (c or "").strip() for c in row):
                    continue
                while len(row) <= max_idx:
                    row.append("")
                prov = _normalize_cell(row[CSV_COL_PROVINCE])
                ward = _normalize_cell(row[CSV_COL_WARD])
                if not prov or not ward:
                    continue
                by_province[prov].add(ward)
                total_rows += 1
    except OSError as e:
        raise OSError(f"Không đọc được CSV {csv_path}: {e}") from e
    except UnicodeDecodeError as e:
        raise ValueError(
            f"Lỗi encoding khi đọc CSV {csv_path}: {e}. Dùng UTF-8 hoặc UTF-8 BOM."
        ) from e

    provinces_sorted = _sort_vi(list(by_province.keys()))
    data: Dict[str, List[str]] = {}
    ward_total = 0
    for p in provinces_sorted:
        wards = _sort_vi(list(by_province[p]))
        data[p] = wards
        ward_total += len(wards)

    out = {
        "generated_from": base,
        "source_format": "csv",
        "province_column": "C",
        "ward_column": "I",
        "header_row": 1,
        "data_start_row": 2,
        "items_count": ward_total,
        "provinces_count": len(data),
        "data": data,
    }
    return out, data


def _import_from_excel(excel_path: str) -> Tuple[dict, Dict[str, List[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        print("ERROR: đọc Excel cần openpyxl: pip install openpyxl", file=sys.stderr)
        raise SystemExit(1) from e

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Không có sheet '{SHEET_NAME}'. Các sheet có sẵn: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    by_province: Dict[str, set] = defaultdict(set)
    total_rows = 0

    for row_idx in range(XLS_DATA_START_ROW, (ws.max_row or XLS_DATA_START_ROW) + 1):
        prov = _normalize_cell(ws.cell(row=row_idx, column=XLS_COL_PROVINCE).value)
        ward = _normalize_cell(ws.cell(row=row_idx, column=XLS_COL_WARD).value)
        if not prov and not ward:
            continue
        if not prov or not ward:
            continue
        by_province[prov].add(ward)
        total_rows += 1

    wb.close()

    provinces_sorted = _sort_vi(list(by_province.keys()))
    data: Dict[str, List[str]] = {}
    ward_total = 0
    for p in provinces_sorted:
        wards = _sort_vi(list(by_province[p]))
        data[p] = wards
        ward_total += len(wards)

    out = {
        "generated_from": os.path.basename(excel_path),
        "source_format": "xlsx",
        "province_column": "D",
        "ward_column": "J",
        "sheet": SHEET_NAME,
        "header_row": XLS_HEADER_ROW,
        "data_start_row": XLS_DATA_START_ROW,
        "items_count": ward_total,
        "provinces_count": len(data),
        "data": data,
    }
    return out, data


def run_import(
    input_path: Optional[str] = None,
    output_path: Optional[str] = None,
) -> Tuple[dict, Dict[str, List[str]]]:
    """
    Đọc CSV (mặc định) hoặc Excel nếu đuôi .xlsx.
    Ghi `static/location_master.json`.
    """
    resolved = _resolve_input_path(input_path)
    output_path = output_path or default_output_path()

    if resolved.lower().endswith(".csv"):
        out, data = _import_from_csv(resolved)
    elif resolved.lower().endswith((".xlsx", ".xlsm")):
        out, data = _import_from_excel(resolved)
    else:
        raise ValueError(f"Định dạng không hỗ trợ: {resolved} (chỉ .csv hoặc .xlsx)")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    return out, data


def main():
    ap = argparse.ArgumentParser(description="Import CSV/Excel → location_master.json")
    ap.add_argument(
        "--input",
        "-i",
        default=None,
        help="Đường dẫn file CSV hoặc Excel (mặc định: data/Danh-muc-Phuong-xa_moi.csv)",
    )
    ap.add_argument("--out", default=None, help="Đường dẫn JSON đầu ra")
    args = ap.parse_args()
    try:
        meta, data = run_import(args.input, args.out)
    except Exception as e:
        print(f"ERROR: {type(e).__name__}: {e}", file=sys.stderr)
        raise SystemExit(1) from e

    print("")
    print("=== Location master import OK ===")
    print(f"Nguồn: {meta.get('generated_from')} ({meta.get('source_format', '?')})")
    print(f"Số tỉnh/thành phố: {meta['provinces_count']}")
    print(f"Số phường/xã (tổng bản ghi sau dedupe): {meta['items_count']}")
    print(f"Đã ghi: {default_output_path() if not args.out else args.out}")
    provs = _sort_vi(list(data.keys()))
    print("5 tỉnh/thành phố đầu (theo thứ tự đã sort):", ", ".join(provs[:5]) or "(trống)")

    def sample_wards(name: str, label: str):
        if name not in data:
            print(f"5 phường/xã mẫu {label}: (không có trong master)")
            return
        w = data[name][:5]
        print(f"5 phường/xã mẫu {label}: {', '.join(w)}")

    sample_wards("Thành phố Hà Nội", "Hà Nội")
    hcm_key = next((k for k in data if "Hồ Chí Minh" in k), None)
    if hcm_key:
        sample_wards(hcm_key, "TP.HCM")
    else:
        print("5 phường/xã mẫu TP.HCM: (không có khóa tỉnh/TP chứa 'Hồ Chí Minh' trong master)")
    print("")


if __name__ == "__main__":
    main()
